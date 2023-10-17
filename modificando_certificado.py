from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt
import pyautogui


caminho_arquivo = "assets/Lutadores.xlsx"
planilha_lutadores = load_workbook(filename=caminho_arquivo)
planilha_selecionada = planilha_lutadores['Dados']

for i in range(2, len(planilha_selecionada['A']) + 1):  # De 2 porque o A1, B1, ..., são, normalmente, os títulos.
    nome = planilha_selecionada[f'A{i}'].value
    peso = float(planilha_selecionada[f'C{i}'].value)

    if peso >= 90:
        categoria = 'pesados'
    elif peso >= 70:
        categoria = 'médios'
    else:
        categoria = 'leves'

    certificado = Document("assets/certificado_base.docx")

    for paragrafo in certificado.paragraphs:
        if '@nome' in paragrafo.text:
            paragrafo.text = paragrafo.text.replace('@nome', nome).replace('@categoria', categoria)
            fonte = certificado.styles["Normal"].font
            fonte.size = Pt(18)
            fonte.name = "Cambria"

    caminho_certificado = f'assets/certificado_{nome}.docx'

    certificado.save(caminho_certificado)